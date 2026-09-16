#!/usr/bin/env python3
"""
QASE Schema Validator
======================

Validates a QASE "Test Designer / repository import" .xlsx file BEFORE it
gets uploaded to QASE, instead of relying on a manual crosscheck to catch
empty or inconsistent fields after the fact.

It checks, per row:
  - required fields are not empty (configurable list)
  - enum fields have valid values (configurable per field)
  - the title matches an expected pattern and its embedded number matches
    the row's id
  - ids are unique (duplicates are always flagged; gaps are a warning,
    since a batch of new cases appended to an existing suite legitimately
    starts mid-sequence)
  - the internal numbering of steps_actions / steps_result / steps_data is
    well-formed (no gaps, no duplicates) in each column independently

Design note: differing step counts BETWEEN steps_actions and steps_result
are reported only as a warning, not an error. Real-world suites commonly
have a single action step open into several separately-numbered result
assertions (or the reverse) — that's a valid authoring pattern, not a
defect. Treating it as an error produces false positives on suites that
are already correct.

Nothing here is tied to a specific project or company: all valid values,
the title pattern, and which fields are required are passed through a
config dict/JSON — see DEFAULT_CONFIG below, and examples/config.example.json
for how to point it at a different naming convention.

Usage:
    python3 qase_schema_validator.py file.xlsx [--config config.json] [--out report.md]

Exit code is 1 if any error was found, 0 otherwise (errors vs. warnings
are always listed separately) — convenient for a pre-import check in a
script or CI step.
"""

import sys
import re
import json
import argparse
from collections import defaultdict, Counter

try:
    import openpyxl
except ImportError:
    sys.exit("Missing openpyxl. Install with: pip install openpyxl")


DEFAULT_CONFIG = {
    # Fields that must not be empty in any row.
    "required_fields": [
        "v2.id", "title", "description", "preconditions", "tags",
        "priority", "severity", "type", "behavior", "automation",
        "status", "is_flaky", "layer", "steps_type",
        "steps_actions", "steps_result", "suite_id", "suite", "is_muted",
    ],
    # Known-valid values per field. A field not listed here is treated as
    # free text and its value is never checked against an enum.
    "enum_values": {
        "priority": ["High", "Medium", "Low"],
        "severity": ["Critical", "Normal", "Minor"],
        "type": ["Functional", "Integration"],
        "behavior": ["Positive", "Negative"],
        "automation": ["To Be Automated", "Automated", "Not To Automate"],
        "status": ["Draft", "Actual"],
        "is_flaky": ["False", "True"],
        "is_muted": ["False", "True"],
        "layer": ["Backend", "Frontend"],
        "steps_type": ["Classic"],
    },
    # Expected title pattern. Group 1 must be the case number, and must
    # match the id field (adjust the prefix/padding to your own convention).
    "title_pattern": r"^TC(\d{3}) - .+",
    "title_id_field": "v2.id",
    # Fields that must hold the SAME value across every row of the file
    # (a single suite shouldn't mix two suite_ids, or two slightly
    # different suite path strings because of a typo).
    "suite_consistency_fields": ["suite_id", "suite_parent_id", "suite"],
    # Field used to check uniqueness/sequencing.
    "id_field": "v2.id",
}


def load_config(path):
    cfg = json.loads(json.dumps(DEFAULT_CONFIG))  # deep copy
    if path:
        with open(path, encoding="utf-8") as f:
            user_cfg = json.load(f)
        cfg.update(user_cfg)
    return cfg


def read_rows(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    col_idx = {h: i for i, h in enumerate(headers) if h}
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        row = {h: r[i] for h, i in col_idx.items()}
        if row.get("v2.id") is None and not any(v not in (None, "") for v in row.values()):
            continue  # fully empty row
        rows.append(row)
    return headers, col_idx, rows


def is_empty(v):
    return v is None or (isinstance(v, str) and v.strip() == "")


def numbered_items(text):
    """Extract the leading numbers from a block like '1. ...\n2. ...'."""
    if not text:
        return []
    return [int(m.group(1)) for m in re.finditer(r"(?:^|\n)\s*(\d+)\s*[\.\)]", str(text))]


def check_sequential(nums):
    """True if nums is 1..N with no gaps or duplicates (order doesn't matter)."""
    if not nums:
        return True, "empty"
    s = sorted(nums)
    expected = list(range(1, len(set(s)) + 1))
    if len(s) != len(set(s)):
        dupes = [n for n, c in Counter(s).items() if c > 1]
        return False, f"duplicate numbers: {dupes}"
    if s != expected:
        return False, f"non-sequential numbering (expected 1..{len(s)}, found {s})"
    return True, "ok"


def validate(headers, col_idx, rows, cfg):
    errors = []   # (case_ref, message)
    warnings = []

    expected_cols = set(cfg["required_fields"]) | set(cfg["enum_values"].keys())
    missing_cols = [c for c in expected_cols if c not in col_idx]
    if missing_cols:
        errors.append((
            "FILE",
            "The file doesn't have the expected QASE 'Test Designer' schema — "
            f"missing columns: {sorted(missing_cols)}. "
            "Is this a Test Run / execution export instead of a repository import file?"
        ))
        return errors, warnings  # no point checking row by row without the base columns

    if not rows:
        errors.append(("FILE", "The file has no data rows."))
        return errors, warnings

    def ref(row):
        vid = row.get(cfg["id_field"])
        title = row.get("title") or "(no title)"
        return f"id={vid} — {title}"

    # 1. Required fields not empty
    for row in rows:
        for field in cfg["required_fields"]:
            if is_empty(row.get(field)):
                errors.append((ref(row), f"required field is empty: '{field}'"))

    # 2. Valid enum values
    for row in rows:
        for field, valid in cfg["enum_values"].items():
            val = row.get(field)
            if is_empty(val):
                continue  # already reported above if it was required
            if str(val) not in valid:
                errors.append((ref(row), f"invalid value in '{field}': '{val}' (valid: {valid})"))

    # 3. Title format + correlation with the id
    pat = re.compile(cfg["title_pattern"])
    for row in rows:
        title = row.get("title")
        vid = row.get(cfg["title_id_field"])
        if is_empty(title):
            continue  # already reported
        m = pat.match(str(title))
        if not m:
            errors.append((ref(row), f"title doesn't match the expected pattern ({cfg['title_pattern']}): '{title}'"))
            continue
        try:
            title_num = int(m.group(1))
            if vid is not None and int(float(vid)) != title_num:
                errors.append((ref(row), f"the number in the title ({m.group(1)}) doesn't match {cfg['title_id_field']}={vid}"))
        except (ValueError, TypeError):
            pass

    # 4. Unique, sequential ids
    ids = []
    for row in rows:
        vid = row.get(cfg["id_field"])
        if vid is None:
            continue
        try:
            ids.append(int(float(vid)))
        except (ValueError, TypeError):
            errors.append((ref(row), f"'{cfg['id_field']}' is not numeric: {vid!r}"))
    dupes = [n for n, c in Counter(ids).items() if c > 1]
    if dupes:
        errors.append(("FILE", f"duplicate ids in '{cfg['id_field']}': {sorted(dupes)}"))
    if ids:
        s = sorted(set(ids))
        gaps = [n for n in range(s[0], s[-1] + 1) if n not in s]
        if gaps:
            warnings.append(("FILE", f"gaps in '{cfg['id_field']}' numbering: {gaps} "
                                       "(expected if these are new cases appended to an existing "
                                       "suite — confirm it wasn't unintentional)"))

    # 5. Steps: well-formed internal numbering in actions and result;
    #    neither one empty. A DIFFERENT count between actions and result is
    #    not an error — it's a valid, common pattern (one action step can
    #    open into several separately-numbered result assertions).
    for row in rows:
        actions = row.get("steps_actions")
        result = row.get("steps_result")
        if is_empty(actions) or is_empty(result):
            continue  # already reported as a required-field gap
        ok_a, msg_a = check_sequential(numbered_items(actions))
        if not ok_a:
            errors.append((ref(row), f"'steps_actions' numbering is malformed: {msg_a}"))
        ok_r, msg_r = check_sequential(numbered_items(result))
        if not ok_r:
            errors.append((ref(row), f"'steps_result' numbering is malformed: {msg_r}"))
        na, nr = len(numbered_items(actions)), len(numbered_items(result))
        if ok_a and ok_r and na != nr:
            warnings.append((ref(row), f"'steps_actions' has {na} step(s) and 'steps_result' has {nr} — "
                                         "not blocking, but worth confirming nothing is missing."))
        data = row.get("steps_data")
        if not is_empty(data):
            ok_d, msg_d = check_sequential(numbered_items(data))
            if not ok_d:
                errors.append((ref(row), f"'steps_data' numbering is malformed: {msg_d}"))

    # 6. Suite/milestone consistency across the whole file
    for field in cfg["suite_consistency_fields"]:
        vals = set(row.get(field) for row in rows if not is_empty(row.get(field)))
        if len(vals) > 1:
            errors.append(("FILE", f"inconsistent values for '{field}' within the same file: {sorted(str(v) for v in vals)} "
                                     "— check for typos; this is exactly the kind of mistake a copy-pasted suite path invites."))

    return errors, warnings


def format_report(xlsx_path, rows, errors, warnings):
    lines = []
    lines.append("# QASE schema validation report")
    lines.append("")
    lines.append(f"File: `{xlsx_path}`")
    lines.append(f"Cases analyzed: {len(rows)}")
    lines.append(f"Errors: {len(errors)} — Warnings: {len(warnings)}")
    lines.append("")
    if not errors and not warnings:
        lines.append("No findings. The file matches the expected schema.")
        return "\n".join(lines)

    if errors:
        lines.append("## Errors (block the QASE import)")
        lines.append("")
        by_case = defaultdict(list)
        for case, msg in errors:
            by_case[case].append(msg)
        for case, msgs in by_case.items():
            lines.append(f"- **{case}**")
            for m in msgs:
                lines.append(f"  - {m}")
        lines.append("")

    if warnings:
        lines.append("## Warnings (non-blocking, use judgment)")
        lines.append("")
        by_case = defaultdict(list)
        for case, msg in warnings:
            by_case[case].append(msg)
        for case, msgs in by_case.items():
            lines.append(f"- **{case}**")
            for m in msgs:
                lines.append(f"  - {m}")
        lines.append("")

    return "\n".join(lines)


def main():
    ap = argparse.ArgumentParser(description="QASE schema validator")
    ap.add_argument("xlsx", help="A .xlsx file in QASE 'Test Designer' import format")
    ap.add_argument("--config", help="Config JSON (overrides DEFAULT_CONFIG)", default=None)
    ap.add_argument("--out", help="Output path for the .md report (default: stdout)", default=None)
    args = ap.parse_args()

    cfg = load_config(args.config)
    headers, col_idx, rows = read_rows(args.xlsx)
    errors, warnings = validate(headers, col_idx, rows, cfg)
    report = format_report(args.xlsx, rows, errors, warnings)

    if args.out:
        with open(args.out, "w", encoding="utf-8") as f:
            f.write(report + "\n")
        print(f"Report written to {args.out}")
    else:
        print(report)

    print(f"\n{len(errors)} error(s), {len(warnings)} warning(s).", file=sys.stderr)
    sys.exit(1 if errors else 0)


if __name__ == "__main__":
    main()
